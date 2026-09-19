# -*- coding: utf-8 -*-
"""Turn a built 16:9 deck into a .pptx, one slide per slide.

    py tools/deck-to-pptx.py <lesson.html> [--lang de] [-o out.pptx]

HOUSE-STYLE §9 already covers the manual route (Ctrl-P, Save as PDF,
Landscape, margins None, background graphics ON) and that is still the right
answer when somebody wants a PDF. This is for the other request, which keeps
coming: a PowerPoint file to send to a person who will not open a browser,
open in whatever they already use, with no live page behind it.

Each slide is a full-bleed picture on a 13.333 x 7.5in blank layout. That is
deliberate rather than lazy: the deck's type, plates and framed artwork are
CSS, and anything that tried to rebuild them as PowerPoint shapes would be a
second implementation of the house style that drifts from the first one.

Two things it does that a screenshot loop does not:

  * **it drives the deck's own Next button.** The match grid, the sort chips
    and the sentence-order strip are built by the engine the first time a
    slide is shown, so a slide reached by toggling `.is-active` by hand
    photographs empty. This cost an afternoon once; do not "simplify" it.
  * **it hides the navigation chrome** - the same set the print stylesheet
    hides. A slide someone has been sent is not navigable, so the language
    picker, the progress bar and the arrows are noise.

The interactive rounds are exported UNANSWERED, which is what you want for
teaching off the file. There is no answer key in the notes: who the file is
being sent to is not knowable from here, and a key in the speaker notes of a
student's copy is worse than no file.

Needs node with playwright (the same dependency check-lesson.js has) and
python-pptx.
"""
import argparse
import glob
import os
import shutil
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

SHOOTER = r"""
const { chromium } = require('playwright');
const path = require('path');
const fs = require('fs');
(async () => {
  const [file, outDir, lang] = process.argv.slice(2);
  fs.mkdirSync(outDir, { recursive: true });
  const pinned = '/opt/pw-browsers/chromium';
  const browser = await chromium.launch(
    fs.existsSync(pinned) ? { executablePath: pinned } : {});
  const page = await browser.newPage({
    viewport: { width: 1280, height: 720 },
    deviceScaleFactor: 1.5,
  });
  await page.goto('file://' + path.resolve(file));
  await page.waitForTimeout(2500);
  await page.evaluate(() => document.fonts.ready);
  if (lang && lang !== 'en') {
    await page.selectOption('#langSelect', lang);
    await page.waitForTimeout(800);
  }
  await page.addStyleTag({ content:
    '.deck-bar, .lang-select, .nav-btn, .deck-rail { display: none !important; }' });
  const total = await page.evaluate(
    () => document.querySelectorAll('.slide[data-type]').length);
  for (let i = 0; i < total; i++) {
    if (i > 0) {
      await page.evaluate(() => {
        const b = document.querySelector('[data-action="next"]:not([hidden])')
               || document.querySelector('.nav-btn[data-dir="1"], #next');
        if (b) b.click();
      });
      await page.waitForTimeout(650);
    }
    await page.locator('.stage').screenshot({
      path: path.join(outDir, String(i + 1).padStart(2, '0') + '.jpg'),
      type: 'jpeg', quality: 86,
    });
  }
  await browser.close();
  console.log(total);
})();
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('deck')
    ap.add_argument('--lang', default='en',
                    help='export in this language (must be complete in the deck)')
    ap.add_argument('-o', '--out')
    a = ap.parse_args()

    try:
        from pptx import Presentation
        from pptx.util import Inches, Emu
    except ImportError:
        sys.exit('needs python-pptx:  py -m pip install python-pptx')

    deck = os.path.abspath(a.deck)
    if not os.path.exists(deck):
        sys.exit('no such deck: %s' % a.deck)
    stem = os.path.splitext(os.path.basename(deck))[0]
    out = a.out or os.path.join(ROOT, '%s%s.pptx' % (
        stem, '' if a.lang == 'en' else '-' + a.lang))

    tmp = tempfile.mkdtemp(prefix='deck-pptx-')
    try:
        # node resolves playwright from the repo's node_modules, so the script
        # has to live inside the repo rather than in the system temp folder.
        js = os.path.join(ROOT, '_deck-to-pptx-%d.js' % os.getpid())
        open(js, 'w', encoding='utf-8', newline='\n').write(SHOOTER)
        try:
            r = subprocess.run(['node', js, deck, tmp, a.lang],
                               cwd=ROOT, capture_output=True, text=True)
        finally:
            os.remove(js)
        if r.returncode:
            sys.exit('render failed:\n' + (r.stderr or r.stdout))

        shots = sorted(glob.glob(os.path.join(tmp, '*.jpg')))
        if not shots:
            sys.exit('render produced no slides')

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        for f in shots:
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(f, Emu(0), Emu(0),
                                 width=prs.slide_width, height=prs.slide_height)
        prs.core_properties.author = 'Forbes English'
        prs.core_properties.comments = 'forbesenglish.com/%s.html' % stem
        prs.save(out)
        print('%s  %d slides  %.1f MB'
              % (out, len(shots), os.path.getsize(out) / 1e6))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    main()
